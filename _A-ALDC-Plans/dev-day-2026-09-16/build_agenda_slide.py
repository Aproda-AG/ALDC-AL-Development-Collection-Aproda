"""Agenda-Folie separat, mit grossem A-ALDC-Logo links (Test-Iteration).
Run: C:/Python314/python.exe build_agenda_slide.py
Output: agenda-test.pptx (nur 1 Folie, zur Review vor Uebernahme ins Hauptdeck)
"""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor

SOURCE_TEMPLATE = "aproda-vorlage.pptx"
LOGO = "a-aldc-logo.png"
OUTPUT = "agenda-test.pptx"

ACCENT = RGBColor(0x3C, 0x6C, 0xE8)  # dk2 aus Aproda-Theme

prs = Presentation(SOURCE_TEMPLATE)


def delete_slide(prs, index):
    """Remove slide + drop its relationship so the part is excluded on save."""
    sld_id = prs.slides._sldIdLst[index]
    prs.part.drop_rel(sld_id.rId)
    prs.slides._sldIdLst.remove(sld_id)


for _ in range(len(prs.slides._sldIdLst)):
    delete_slide(prs, 0)

layouts = {layout.name: layout for layout in prs.slide_layouts}

slide = prs.slides.add_slide(layouts["Nur Titel"])
slide.shapes.title.text = "Agenda"

# grosses Logo links
slide.shapes.add_picture(LOGO, Inches(0.7), Inches(1.7), height=Inches(5.2))

# Agenda-Punkte rechts als Textbox
box = slide.shapes.add_textbox(Inches(6.8), Inches(2.0), Inches(6.0), Inches(4.5))
tf = box.text_frame
tf.word_wrap = True

items = [
    "Von der Theorie zur Praxis: A-ALDC im Ueberblick",
    "Aufbau & Prozess von A-ALDC",
    "Agents im Ueberblick",
    "BCQuality: was & warum",
]
for i, text in enumerate(items):
    p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
    p.text = f"{i + 1}.  {text}"
    p.font.size = Pt(24)
    p.font.color.rgb = ACCENT if i == 0 else RGBColor(0x00, 0x00, 0x00)
    p.space_after = Pt(18)

slide.notes_slide.notes_text_frame.text = "Agenda-Variante mit grossem Logo links, Punkte rechts."

prs.save(OUTPUT)
print(f"Saved {OUTPUT}")
