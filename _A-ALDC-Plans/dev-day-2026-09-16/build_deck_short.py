"""Kurze Draft-Version (8 Folien) fuer den Dev Day: A-ALDC + BCQuality.
Baut auf der echten Aproda-Vorlage auf (Layouts/Master/Branding), inkl. A-ALDC-Logo.
Run: C:/Python314/python.exe build_deck_short.py
Output: A-ALDC-DevDay-kurz.pptx
"""
from pptx import Presentation
from pptx.util import Inches

SOURCE_TEMPLATE = "aproda-vorlage.pptx"
LOGO = "a-aldc-logo.png"
OUTPUT = "A-ALDC-DevDay-kurz.pptx"

prs = Presentation(SOURCE_TEMPLATE)


def delete_slide(prs, index):
    """Remove slide + drop its relationship so the part is excluded on save."""
    sld_id = prs.slides._sldIdLst[index]
    prs.part.drop_rel(sld_id.rId)
    prs.slides._sldIdLst.remove(sld_id)


# nur Layouts/Master behalten, mitgelieferte Beispiel-Folien entfernen
for _ in range(len(prs.slides._sldIdLst)):
    delete_slide(prs, 0)

layouts = {layout.name: layout for layout in prs.slide_layouts}


def set_placeholder_text(slide, idx, lines):
    ph = slide.placeholders[idx]
    tf = ph.text_frame
    tf.clear()
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = line


def add_slide(layout_name, title=None, notes=None):
    slide = prs.slides.add_slide(layouts[layout_name])
    if title is not None and slide.shapes.title is not None:
        slide.shapes.title.text = title
    if notes:
        slide.notes_slide.notes_text_frame.text = notes
    return slide


# 1) Titelfolie
s1 = add_slide(
    "Titelfolie",
    title="A-ALDC",
    notes=(
        "Uebergang: 'Das war die Theorie - grosse Teile dieser Prinzipien "
        "finden wir in A-ALDC, einem Agentic Engineering Toolkit fuer BC AL.'"
    ),
)
set_placeholder_text(s1, 10, ["Agentic Engineering Toolkit fuer Business Central AL"])
s1.shapes.add_picture(LOGO, Inches(11.0), Inches(0.35), height=Inches(1.6))

# 2) Agenda
s2 = add_slide("Agenda", notes="Ueberblick ueber die 4 Blöcke des Teils.")
for idx, text in zip(
    (10, 11, 12, 13),
    [
        "Von der Theorie zur Praxis: A-ALDC im Ueberblick",
        "Aufbau & Prozess von A-ALDC",
        "Agents im Ueberblick",
        "BCQuality: was & warum",
    ],
):
    set_placeholder_text(s2, idx, [text])

# 3) Abschnitt: Theorie -> Praxis
s3 = add_slide(
    "Abschnitt 1",
    title="Von der Theorie zur Praxis",
    notes="Bezug zu Harnisch/Skills-Vortrag herstellen.",
)
set_placeholder_text(
    s3, 1, ["Wo A-ALDC die vorgestellten Prinzipien (Harnisch, Skills, Agents) konkret umsetzt"]
)

# 4) Aufbau von A-ALDC
s4 = add_slide("Titel und Inhalt", title="Aufbau von A-ALDC", notes="Kernbausteine kurz vorstellen.")
set_placeholder_text(
    s4,
    1,
    [
        "Agents – strategische & taktische Rollen (Architect, Developer, Conductor, Presales, Subagents)",
        "Skills – domaenenspezifisches Wissen, wird bei Bedarf geladen",
        "Workflows – wiederholbare Ablaeufe (spec.create, build, pr-prepare, ...)",
        "Instructions – greifen automatisch je Dateityp (z. B. *.al)",
        "Plans – strukturierte Ablage unter .github/plans (Spec, Architektur, Review)",
    ],
)

# 5) Agents im Ueberblick
s5 = add_slide("Titel und Inhalt", title="Agents im Ueberblick", notes="Kurz Rolle je Agent nennen.")
set_placeholder_text(
    s5,
    1,
    [
        "al-architect – Architektur, Datenmodell, Integrationsstrategie",
        "al-developer – Implementierung, Debugging, Fixes",
        "al-conductor – TDD-Orchestrierung (Plan -> Implement -> Review)",
        "al-presales – Schaetzung, SWOT, Projektangebote",
        "+ Subagents fuer Planning / Implementation / Review",
    ],
)

# 6) Abschnitt: BCQuality
s6 = add_slide("Abschnitt 2", title="BCQuality", notes="Ueberleitung zum zweiten Themenblock.")
set_placeholder_text(s6, 1, ["Eine kuratierte, zitierfaehige Wissensbasis fuer Business Central"])

# 7) BCQuality: was & warum
s7 = add_slide(
    "Titel und Inhalt", title="BCQuality: was & warum", notes="Warum das ein Audit-Layer ist, kein Ersatz."
)
set_placeholder_text(
    s7,
    1,
    [
        "Kuratierte Wissensbasis mit zitierbaren Guidance-Dateien",
        "Wird als zweite Workspace-Root eingebunden – kein Einfluss auf die Kompilierung",
        "Audit-/Zitationsschicht, kein Ersatz fuer Instructions & Skills",
        "Review-Subagent konsultiert BCQuality vor der A-G-Checkliste",
        "Findings sind zitierbar & per CI validierbar – keine halluzinierten Quellen",
    ],
)

# 8) Naechste Schritte
s8 = add_slide("Titel und Inhalt", title="Naechste Schritte", notes="Offen fuer Fragen & Feedback.")
set_placeholder_text(
    s8,
    1,
    [
        "Live-Demo im Team einplanen",
        "Feedback & offene Fragen sammeln",
        "Rollout-Fahrplan fuer weitere Projekte",
    ],
)

prs.save(OUTPUT)
print(f"Saved {OUTPUT} ({len(prs.slides)} Folien)")
