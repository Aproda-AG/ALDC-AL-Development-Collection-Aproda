"""Test: reuse the real Aproda .pptx template (layouts, masters, branding).
Run: C:/Python314/python.exe test-pptx-aproda-template.py
Output: test-aproda-template.pptx
"""
from pptx import Presentation
from pptx.util import Pt

SOURCE_TEMPLATE = "aproda-vorlage.pptx"
OUTPUT = "test-aproda-template.pptx"

prs = Presentation(SOURCE_TEMPLATE)


def delete_slide(prs, index):
    """Remove slide + drop its relationship so the part is excluded on save."""
    sld_id = prs.slides._sldIdLst[index]
    prs.part.drop_rel(sld_id.rId)
    prs.slides._sldIdLst.remove(sld_id)


# drop the 8 sample slides that ship with the template, keep only layouts/masters
for _ in range(len(prs.slides._sldIdLst)):
    delete_slide(prs, 0)

layouts = {layout.name: layout for layout in prs.slide_layouts}


def add_slide(layout_name, title=None, body_lines=None, body_idx=1):
    slide = prs.slides.add_slide(layouts[layout_name])
    if title is not None and slide.shapes.title is not None:
        slide.shapes.title.text = title
    if body_lines:
        placeholder = slide.placeholders[body_idx]
        tf = placeholder.text_frame
        tf.clear()
        for i, line in enumerate(body_lines):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.text = line
    return slide


s1 = add_slide("Titelfolie", title="A-ALDC & BCQuality")
s1.notes_slide.notes_text_frame.text = "Test-Slide 1/4 - Pipeline-Check (Aproda-Template)"

add_slide(
    "Agenda",
    body_idx=10,
    body_lines=[
        "Theorie -> Praxis: wo A-ALDC die Prinzipien umsetzt",
        "Aufbau & Prozess von A-ALDC",
        "Agents im Ueberblick",
        "BCQuality: was, warum",
    ],
)

add_slide("Abschnitt 1", title="Platzhalter: Architektur-Diagramm")

add_slide(
    "Titel und Inhalt",
    title="Test Ende",
    body_lines=["Naechster Schritt: Format entscheiden, dann Inhalt fuellen."],
)

prs.save(OUTPUT)
print(f"Saved {OUTPUT}")
