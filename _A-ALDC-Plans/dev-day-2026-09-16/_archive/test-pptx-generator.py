"""Test script: builds a minimal .pptx to validate the python-pptx pipeline.
Run: C:/Python314/python.exe test-pptx-generator.py
Output: test-deck.pptx (same folder)
"""
from pptx import Presentation
from pptx.util import Inches, Pt

OUTPUT = "test-deck.pptx"

SLIDES = [
    {
        "title": "A-ALDC",
        "body": ["Agentic Engineering Toolkit fuer BC AL", "(Test-Slide 1/4 - Pipeline-Check)"],
        "notes": "Test-Deck um die python-pptx Pipeline zu pruefen. Kein finaler Inhalt.",
    },
    {
        "title": "Agenda (Platzhalter)",
        "body": [
            "Theorie -> Praxis: wo A-ALDC die Prinzipien umsetzt",
            "Aufbau & Prozess von A-ALDC",
            "Agents im Ueberblick",
            "BCQuality: was, warum",
        ],
        "notes": "Kurzer Uebergang von den allgemeinen Prinzipien (Harnisch, Skills) zu A-ALDC.",
    },
    {
        "title": "Platzhalter: Architektur-Diagramm",
        "body": ["[Hier wuerde ein Mermaid-Diagramm als Bild eingefuegt werden]"],
        "notes": "Diagramm zeigt Agents/Skills/Workflows-Zusammenspiel.",
    },
    {
        "title": "Test Ende",
        "body": ["Naechster Schritt: Format entscheiden, dann Inhalt fuellen."],
        "notes": "Abschluss-Slide des Tests.",
    },
]

prs = Presentation()
layout = prs.slide_layouts[1]  # Title and Content

for slide_data in SLIDES:
    slide = prs.slides.add_slide(layout)
    slide.shapes.title.text = slide_data["title"]
    body_placeholder = slide.placeholders[1]
    tf = body_placeholder.text_frame
    tf.clear()
    for i, line in enumerate(slide_data["body"]):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = line
        p.font.size = Pt(24)
    slide.notes_slide.notes_text_frame.text = slide_data["notes"]

prs.save(OUTPUT)
print(f"Saved {OUTPUT}")
